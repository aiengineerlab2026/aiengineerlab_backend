"""My Resume lab: ranking maths, BM25 preprocessing, the three retrieval methods and the HTTP API (no network, fake embedder)."""

from __future__ import annotations

import hashlib

import numpy as np
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.config import Settings
from app.error_handlers import register_error_handlers
from app.errors import ExternalServiceError, ValidationFailed
from app.myresume import presentation
from app.myresume.application.evaluation import RankingEvaluationService, skill_coverage
from app.myresume.container import MyResumeInfra
from app.myresume.domain.ranking import match_percent, rrf_fuse
from app.myresume.domain.retrieval import Bm25Settings, Method, RetrievalConfig
from app.myresume.infrastructure.bm25 import Bm25Registry
from app.myresume.infrastructure.data.jobs import JOBS, JOBS_BY_ID
from app.rate_limit import SlidingWindowLimiter
from app.restaurant.infrastructure.llm.quota import QuotaTracker

# ----------------------------------------------------------------------------------------------- ranking maths


def test_rrf_scores_and_order():
    fused = dict(rrf_fuse([["a", "b", "c"], ["b", "a", "d"]], k=60))
    assert fused["a"] == pytest.approx(1 / 61 + 1 / 62)
    assert fused["b"] == pytest.approx(1 / 62 + 1 / 61)
    assert fused["c"] == pytest.approx(1 / 63) and fused["d"] == pytest.approx(1 / 63)
    # tie between a and b keeps first appearance; c before d likewise
    assert [j for j, _ in rrf_fuse([["a", "b", "c"], ["b", "a", "d"]])] == ["a", "b", "c", "d"]


def test_rrf_item_in_both_lists_beats_item_in_one():
    order = [j for j, _ in rrf_fuse([["x", "shared"], ["y", "shared"]], k=60)]
    assert order[0] == "shared"


def test_rrf_a_list_votes_once_per_item_and_k_is_validated():
    assert dict(rrf_fuse([["a", "a"]], k=1))["a"] == pytest.approx(1 / 2)
    with pytest.raises(ValueError):
        rrf_fuse([["a"]], k=0)


def test_match_percent_per_method():
    assert match_percent("vector", 0.8123, 0.9) == 81
    top, low = 2 / 61, 1 / 80  # best fused score, the last one returned
    assert match_percent("hybrid", top, top, low) == 100 and match_percent("hybrid", low, top, low) == 0
    assert match_percent("hybrid", (top + low) / 2, top, low) == 50  # min-max over the returned jobs
    assert match_percent("hybrid", 0.03, 0.03, 0.03) == 100  # one job (or all tied): nothing to spread
    assert match_percent("bm25", 5.0, 10.0) == 50 and match_percent("bm25", 1.0, 0.0) == 0
    assert match_percent("vector", 1.7, 1.0) == 100  # always within 0-100


# ----------------------------------------------------------------------------------------------- BM25 / spaCy


def test_bm25_preprocessing_settings_change_the_tokens():
    reg = Bm25Registry(JOBS)
    resume = "The developers were building APIs, and testing them!"
    full = reg.get(Bm25Settings())
    raw = reg.get(Bm25Settings(remove_stopwords=False, lemmatize=False, lowercase=False, remove_punctuation=False, min_token_length=1))
    toks = full._pipe.tokens([resume], full.settings)[0]
    raw_toks = raw._pipe.tokens([resume], raw.settings)[0]
    assert "the" not in toks and "!" not in toks and "build" in toks  # stop words and punctuation gone, lemmatised
    assert "The" in raw_toks and "!" in raw_toks and "building" in raw_toks


def test_bm25_ranks_the_matching_job_first():
    scores, terms = Bm25Registry(JOBS).get(Bm25Settings()).score(
        "I develop iPhone apps in Swift and SwiftUI and publish to the App Store"
    )
    assert terms > 0
    assert JOBS[int(np.argmax(scores))].id == "ios-engineer"


def test_bm25_with_no_usable_tokens_scores_zero():
    scores, terms = Bm25Registry(JOBS).get(Bm25Settings()).score("the of and ...")
    assert terms == 0 and not scores.any()


# ----------------------------------------------------------------------------------------------- retrieval with a fake embedder


class FakeEmbedder:
    """Deterministic bag-of-hashed-words embedder: similar texts get similar vectors, so cosine search is meaningful."""

    dims = 64

    def _vec(self, text: str) -> list[float]:
        v = np.zeros(self.dims)
        for w in text.lower().split():
            w = w.strip(".,;:()!?")
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dims] += 1
        n = np.linalg.norm(v)
        return list(v / n if n else v)

    async def embed_documents(self, texts):
        return [self._vec(t) for t in texts]

    async def embed_query(self, text):
        return self._vec(text)


@pytest.fixture
def infra() -> MyResumeInfra:
    i = MyResumeInfra(Settings(app_env="test", llm_api_key="x"), QuotaTracker())
    fake = FakeEmbedder()
    i.embedder_for = lambda spec: fake  # type: ignore[method-assign]
    i.vectors._embedder_for = i.embedder_for
    i.search._embedder_for = i.embedder_for
    return i


@pytest.fixture
def spec(infra):
    return infra.resolve_embedding(None)


METHODS = (Method.BM25, Method.VECTOR, Method.HYBRID)
RESUME = (
    "Alex Morgan. Software developer with six years of experience building server-side applications. Designed and shipped "
    "HTTP APIs in Python with Django and FastAPI for an e-commerce marketplace, backed by PostgreSQL and Redis. "
    "Containerised services with Docker and ran them on AWS. Wrote automated tests, profiled slow queries and cut average "
    "response time by 40 percent. Mentored two junior developers and took part in on-call rotations."
)


async def test_each_method_returns_top_k_with_its_own_scores(infra, spec):
    out = {}
    for m in METHODS:
        out[m] = await infra.search.match(RESUME, RetrievalConfig(method=m, top_k=4), spec)
        assert len(out[m].results) == 4
        assert [r.rank for r in out[m].results] == [1, 2, 3, 4]
    bm25, vec, hyb = (out[m].results for m in METHODS)
    assert all(r.bm25_score is not None and r.vector_score is None for r in bm25)
    assert all(r.vector_score is not None and r.bm25_score is None for r in vec)
    assert all(r.bm25_score is not None and r.vector_score is not None and r.rrf_score is not None for r in hyb)
    assert [r.score for r in bm25] == sorted([r.score for r in bm25], reverse=True)
    assert out[Method.BM25].query_terms and out[Method.BM25].embed_ms is None and out[Method.HYBRID].embed_ms is not None


async def test_hybrid_is_the_rrf_of_the_two_top_n_lists(infra, spec):
    cfg = RetrievalConfig(method=Method.HYBRID, top_k=5, candidate_count=20, rrf_k=60)
    hybrid = await infra.search.match(RESUME, cfg, spec)
    bm25 = await infra.search.match(RESUME, cfg.model_copy(update={"method": Method.BM25, "top_k": 20}), spec)
    vec = await infra.search.match(RESUME, cfg.model_copy(update={"method": Method.VECTOR, "top_k": 20}), spec)
    expected = rrf_fuse([[r.job_id for r in bm25.results], [r.job_id for r in vec.results]], k=60)[:5]
    assert [r.job_id for r in hybrid.results] == [j for j, _ in expected]
    assert [r.rrf_score for r in hybrid.results] == pytest.approx([s for _, s in expected])


async def test_rrf_constant_changes_scores_not_membership_of_candidates(infra, spec):
    a = await infra.search.match(RESUME, RetrievalConfig(top_k=3, rrf_k=1), spec)
    b = await infra.search.match(RESUME, RetrievalConfig(top_k=3, rrf_k=1000), spec)
    assert a.results[0].rrf_score > b.results[0].rrf_score


async def test_vector_threshold_drops_weak_candidates(infra, spec):
    open_ = await infra.search.match(RESUME, RetrievalConfig(method=Method.VECTOR, top_k=10), spec)
    cut = open_.results[2].vector_score
    strict = await infra.search.match(RESUME, RetrievalConfig(method=Method.VECTOR, top_k=10, vector_threshold=cut), spec)
    assert 0 < len(strict.results) < len(open_.results)
    assert all(r.vector_score >= cut - 1e-9 for r in strict.results)
    none = await infra.search.match(RESUME, RetrievalConfig(method=Method.VECTOR, top_k=5, vector_threshold=1.0), spec)
    assert none.results == [] and none.notices


async def test_short_resume_is_rejected(infra, spec):
    with pytest.raises(ValidationFailed):
        await infra.search.match("too short", RetrievalConfig(), spec)


def test_skill_coverage_counts_whole_words_only():
    job = JOBS_BY_ID["backend-python"]
    assert skill_coverage(job, RESUME) >= 0.6
    assert skill_coverage(job, "I enjoy gardening and cooking for friends on weekends.") == 0.0
    assert skill_coverage(job, "dockerized") == 0.0  # "Docker" inside another word is not a mention


# ----------------------------------------------------------------------------------------------- HTTP API


@pytest.fixture
def client(infra):
    app = FastAPI()
    app.state.myresume = infra
    register_error_handlers(app)
    app.include_router(presentation.router, prefix="/api/v1")
    return TestClient(app)


def test_catalogue_endpoints(client):
    assert len(client.get("/api/v1/myresume/jobs").json()) == len(JOBS)
    d = client.get("/api/v1/myresume/config/defaults").json()
    assert [m["id"] for m in d["methods"]] == ["bm25", "vector", "hybrid"]
    assert d["retrieval"]["rrf_k"] == 60 and d["retrieval"]["candidate_count"] == 20 and d["lemmatizer_available"]


def test_models_offer_embeddings_from_gemini_and_cloudflare(client):
    m = client.get("/api/v1/myresume/models").json()
    assert {p["id"] for p in m["providers"]} == {"google", "cloudflare"}
    assert {o["provider"] for o in m["embedding"]} == {"google", "cloudflare"}
    assert "chat" not in m and m["default_embedding"]["model"]


def test_match_endpoint_and_validation(client):
    r = client.post("/api/v1/myresume/match", json={"resume_text": RESUME, "config": {"method": "hybrid", "top_k": 3}})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["method"] == "hybrid" and len(body["results"]) == 3
    top = body["results"][0]
    assert {"bm25_score", "vector_score", "rrf_score", "match_percent", "job"} <= top.keys() and top["job"]["title"]
    assert all(0 <= r["match_percent"] <= 100 for r in body["results"])
    assert client.post("/api/v1/myresume/match", json={"resume_text": RESUME, "config": {"top_k": 99}}).status_code == 422
    assert client.post("/api/v1/myresume/match", json={"resume_text": "short"}).status_code == 422


def test_unknown_model_is_rejected(client):
    r = client.post(
        "/api/v1/myresume/match",
        json={"resume_text": RESUME, "embedding": {"provider": "cloudflare", "model": "nope"}},
    )
    assert r.status_code == 422


def test_uploads_are_limited_to_three_and_an_unreadable_file_is_not_counted(client):
    url = "/api/v1/myresume/resume"
    assert client.get(f"{url}/usage").json()["limit"] == 3
    assert client.post(f"{url}/extract", files={"file": ("cv.png", b"\x89PNG....")}).status_code == 422
    assert client.get(f"{url}/usage").json()["used"] == 0
    for n in range(1, 4):
        r = client.post(f"{url}/extract", files={"file": ("cv.docx", _docx_bytes())})
        assert r.status_code == 200 and r.json()["usage"]["used"] == n
    fourth = client.post(f"{url}/extract", files={"file": ("cv.docx", _docx_bytes())})
    assert fourth.status_code == 429 and "can be used 3 times" in fourth.text


# ----------------------------------------------------------------------------------------------- resume file upload


def _docx_bytes() -> bytes:
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe   Backend developer")
    doc.add_paragraph("Built REST APIs in Python and FastAPI.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker, AWS"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """A minimal one-page PDF with a text layer (offsets are computed, so the file is well formed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)


def test_extract_text_from_docx_paragraphs_and_tables():
    from app.myresume.infrastructure.resume_files import extract_resume_text

    text = extract_resume_text("cv.docx", _docx_bytes())
    assert "Jane Doe Backend developer" in text  # runs of spaces collapse
    assert "Built REST APIs in Python and FastAPI." in text
    assert "Skills | Docker, AWS" in text


def test_extract_text_from_pdf():
    from app.myresume.infrastructure.resume_files import extract_resume_text

    assert "Senior Python developer with AWS" in extract_resume_text("cv.pdf", _pdf_bytes("Senior Python developer with AWS"))


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("cv.pdf", b"", "empty"),
        ("cv.txt", b"just some text", "Only PDF and DOCX"),
        ("cv.doc", b"PK\x03\x04...", "Only PDF and DOCX"),
        ("cv.pdf", b"%PDF-1.4 garbage", "could not be read"),
        ("cv.docx", b"PK\x03\x04 not a zip", "could not be read"),
    ],
)
def test_extract_rejects_unusable_files(name, data, message):
    from app.myresume.infrastructure.resume_files import extract_resume_text

    with pytest.raises(ValidationFailed, match=message):
        extract_resume_text(name, data)


def test_extract_reports_a_pdf_without_text():
    import io

    from pypdf import PdfWriter

    from app.myresume.infrastructure.resume_files import extract_resume_text

    w, buf = PdfWriter(), io.BytesIO()
    w.add_blank_page(612, 792)  # what a scanned (image-only) PDF looks like to a text extractor
    w.write(buf)
    with pytest.raises(ValidationFailed, match="No text was found"):
        extract_resume_text("scan.pdf", buf.getvalue())


def test_extract_endpoint(client):
    ok = client.post("/api/v1/myresume/resume/extract", files={"file": ("cv.docx", _docx_bytes())})
    assert ok.status_code == 200, ok.text
    body = ok.json()
    assert body["filename"] == "cv.docx" and "FastAPI" in body["text"] and not body["truncated"]
    assert body["characters"] == len(body["text"])
    bad = client.post("/api/v1/myresume/resume/extract", files={"file": ("cv.png", b"\x89PNG....")})
    assert bad.status_code == 422 and "PDF and DOCX" in bad.json()["error"]["message"]
    big = client.post("/api/v1/myresume/resume/extract", files={"file": ("cv.pdf", b"%PDF" + b"0" * (5 * 1024 * 1024))})
    assert big.status_code == 422 and "too large" in big.json()["error"]["message"]


# ----------------------------------------------------------------------------------------------- ranking evaluation (ranx)


async def test_ranking_metrics_are_computed_by_ranx_against_the_resume_judgements():
    svc = RankingEvaluationService(JOBS)
    judged = svc.judgements(RESUME)
    assert judged and judged[0][0] == "backend-python" and all(g in (2, 3) for _, _, g in judged)
    relevant = [j for j, _, _ in judged]
    # a perfect ranking: the relevant jobs first, best coverage first
    k = min(3, len(relevant))
    perfect = await svc.run(RESUME, relevant[:k])
    assert perfect.k == k and perfect.precision == 1.0 and perfect.mrr == 1.0 and perfect.ndcg == pytest.approx(1.0)
    assert perfect.recall == pytest.approx(k / len(relevant))
    assert [r.rank for r in perfect.relevant][:k] == list(range(1, k + 1)) and all(r.rank is None for r in perfect.relevant[k:])
    # the same jobs behind two irrelevant ones: precision drops, MRR is the reciprocal of the first hit's rank
    worse = await svc.run(RESUME, ["accountant", "registered-nurse", relevant[0]])
    assert worse.precision == pytest.approx(1 / 3) and worse.mrr == pytest.approx(1 / 3) and worse.ndcg < 1.0
    nothing = await svc.run(RESUME, ["accountant", "registered-nurse"])
    assert (nothing.precision, nothing.recall, nothing.ndcg, nothing.mrr) == (0.0, 0.0, 0.0, 0.0)


async def test_ndcg_uses_the_grades_while_precision_does_not():
    svc = RankingEvaluationService(JOBS)
    cv = "Python FastAPI Django PostgreSQL REST APIs Docker AWS unit testing Jenkins"
    grades = {j: g for j, _, g in svc.judgements(cv)}
    assert grades["backend-python"] == 3 and grades["devops-engineer"] == 2  # all its skills vs a third of them
    best_first = await svc.run(cv, ["backend-python", "devops-engineer"])
    weak_first = await svc.run(cv, ["devops-engineer", "backend-python"])
    assert best_first.precision == weak_first.precision and best_first.recall == weak_first.recall
    assert best_first.ndcg > weak_first.ndcg


async def test_evaluation_needs_a_known_non_empty_ranking():
    svc = RankingEvaluationService(JOBS)
    with pytest.raises(ValidationFailed):
        await svc.run(RESUME, [])
    with pytest.raises(ValidationFailed):
        await svc.run(RESUME, ["backend-python", "nope"])


async def test_evaluation_without_a_relevant_job_says_metrics_are_undefined():
    svc = RankingEvaluationService(JOBS)
    res = await svc.run("I enjoy gardening, cooking and long walks with my dog on quiet weekends in the countryside.", ["accountant"])
    assert res.relevant == [] and any("not defined" in n for n in res.notices)
    assert None is res.precision is res.recall is res.ndcg is res.mrr


def test_evaluation_endpoint_and_the_removed_endpoints(client):
    url = "/api/v1/myresume"
    top = client.post(f"{url}/match", json={"resume_text": RESUME, "config": {"method": "bm25", "top_k": 5}}).json()["results"]
    ids = [r["job_id"] for r in top]
    r = client.post(f"{url}/evaluation", json={"resume_text": RESUME, "job_ids": ids})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["k"] == 5 and all(0.0 <= body[m] <= 1.0 for m in ("precision", "recall", "ndcg", "mrr"))
    assert body["relevant"][0]["title"] and body["relevant"][0]["company"]
    assert client.post(f"{url}/evaluation", json={"resume_text": RESUME, "job_ids": []}).status_code == 422
    assert client.post(f"{url}/evaluation", json={"resume_text": RESUME, "job_ids": ["nope"]}).status_code == 422
    for gone in ("answer", "explain", "evaluation/answer", "evaluation/answer/llm", "evaluation/run"):
        assert client.post(f"{url}/{gone}", json={}).status_code in (404, 405)


def test_vector_searches_are_limited_to_three_and_only_vector_counts(client, infra):
    url = "/api/v1/myresume"
    search = lambda method: client.post(f"{url}/match", json={"resume_text": RESUME, "config": {"method": method, "top_k": 3}})  # noqa: E731
    assert client.get(f"{url}/match/vector/usage").json()["limit"] == 3
    for n in range(1, 4):
        assert search("vector").status_code == 200
        assert client.get(f"{url}/match/vector/usage").json()["used"] == n
    fourth = search("vector")
    assert fourth.status_code == 429
    message = fourth.json()["error"]["message"]
    assert "Vector search can be used 3 times" in message and "IP" not in message and "day" not in message
    assert search("bm25").status_code == 200 and search("hybrid").status_code == 200  # the other methods are not capped


def test_a_failed_vector_search_is_not_counted(client, infra):
    infra.vector_limiter = SlidingWindowLimiter(1, 1e9)
    url = "/api/v1/myresume"
    bad = {"resume_text": RESUME, "config": {"method": "vector"}, "embedding": {"provider": "cloudflare", "model": "nope"}}
    assert client.post(f"{url}/match", json=bad).status_code == 422  # rejected before it is counted
    assert client.get(f"{url}/match/vector/usage").json()["used"] == 0

    async def broken(*a, **k):
        raise ExternalServiceError("boom")

    infra.search.match = broken  # type: ignore[method-assign]
    assert client.post(f"{url}/match", json={"resume_text": RESUME, "config": {"method": "vector"}}).status_code == 502
    assert client.get(f"{url}/match/vector/usage").json()["used"] == 0  # refunded
